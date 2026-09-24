from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import CellIsRule, FormulaRule, DataBarRule
from openpyxl.chart import BarChart, PieChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.comments import Comment

OUT = "/home/user/mi-bot-trading/plantillas/Plantilla_Libertad_Financiera.xlsx"

NAVY, GOLD, CREAM, INPUT, GRAY = "1F3864", "C9A227", "FBF6E9", "FFF2CC", "F2F2F2"
F = "Arial"
MONEY = '"$"#,##0.00;[Red]-"$"#,##0.00;"-"'
PCT = '0.0%;[Red]-0.0%;"-"'
DATEF = "dd/mm/yyyy"
MONTHF = 'mmmm yyyy'
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

def font(**k):
    k.setdefault("name", F); k.setdefault("size", 10)
    return Font(**k)

def fill(c): return PatternFill("solid", start_color=c, end_color=c)

def title(ws, text, sub=None, width="K"):
    ws.sheet_view.showGridLines = False
    ws["A1"] = text
    ws["A1"].font = font(size=18, bold=True, color="FFFFFF")
    ws.merge_cells(f"A1:{width}1")
    ws["A1"].fill = fill(NAVY)
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 34
    if sub:
        ws["A2"] = sub
        ws["A2"].font = font(italic=True, color="7F6000")
        ws.merge_cells(f"A2:{width}2")
        ws["A2"].fill = fill(CREAM)

def header(ws, row, col, labels):
    for i, l in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=l)
        c.font = font(bold=True, color="FFFFFF")
        c.fill = fill(NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
    ws.row_dimensions[row].height = 30

def style_range(ws, rng, fmt=None, inp=False, calc=False, align=None):
    for row in ws[rng]:
        for c in row:
            c.border = BOX
            c.font = font(color="0000FF") if inp else font()
            if inp: c.fill = fill(INPUT)
            if calc: c.fill = fill(GRAY)
            if fmt: c.number_format = fmt
            if align: c.alignment = Alignment(horizontal=align)

def widths(ws, d):
    for k, v in d.items(): ws.column_dimensions[k].width = v

def name(wb, n, ref):
    wb.defined_names[n] = DefinedName(n, attr_text=ref)

def dv_list(ws, formula, rng, msg=None):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    if msg:
        dv.error = msg; dv.errorTitle = "Valor no válido"
    ws.add_data_validation(dv); dv.add(rng)

wb = Workbook()

# ------------------------------------------------------------------ INSTRUCCIONES
ws = wb.active; ws.title = "Instrucciones"
title(ws, "PLANTILLA DE LIBERTAD FINANCIERA", "Controla tus gastos · Elimina tus deudas · Págate primero · Compra activos", "H")
widths(ws, {"A": 3, "B": 26, "C": 90})
rows = [
    ("CÓMO USARLA", None),
    ("1. Configuración", "Escribe tu nombre, elige si controlarás tus gastos Diario, Semanal o Mensual y ajusta los porcentajes (por defecto la regla de Babilonia 10/20/70). Ajusta tus categorías y su presupuesto mensual."),
    ("2. Ingresos", "Registra cada ingreso: sueldo, ganancia de tu negocio, freelance o ingreso pasivo (rendimientos, rentas, dividendos)."),
    ("3. Gastos", "Anota cada gasto el día que lo haces (o junta tus tickets y captúralos una vez por semana). Todo se resta automáticamente de tu ingreso."),
    ("4. Deudas", "Da de alta cada deuda con su saldo, tasa de interés ANUAL y pago mínimo. La plantilla las ordena sola: primero las que cobran intereses (tarjetas de crédito, la tasa más alta primero) y después los préstamos familiares sin intereses."),
    ("5. Abonos a Deudas", "Cada vez que pagues una deuda, regístralo aquí. El saldo de la deuda baja automáticamente."),
    ("6. Ahorro e Inversión", "Registra lo que te 'pagas primero': fondo de emergencia, inversiones, negocio, educación financiera. Si un activo te genera ingreso, regístralo en Ingresos como 'Ingreso pasivo'."),
    ("7. Resumen", "Tu tablero del mes: cuánto llevas gastado contra tu presupuesto del periodo, semáforo, avance del fondo de emergencia, índice de libertad financiera y fecha estimada para quedar libre de deudas."),
    ("8. Tendencia", "Tu evolución mes a mes durante el año."),
    ("", None),
    ("LEYENDA DE CELDAS", None),
    ("Amarillo / texto azul", "Celdas que TÚ llenas (datos de entrada)."),
    ("Gris / texto negro", "Fórmulas automáticas. No las borres."),
    ("Datos de ejemplo", "La plantilla trae datos de ejemplo de septiembre 2026 para que veas cómo funciona. Bórralos (solo las celdas amarillas) antes de empezar con tus datos."),
    ("", None),
    ("TIPS", None),
    ("Fechas", "Escribe las fechas como dd/mm/aaaa. Las columnas 'Mes' y 'Semana' se calculan solas."),
    ("Filas", "Gastos admite 1,000 registros; Ingresos, Abonos y Ahorro 300 cada uno; Deudas 15; Categorías 20."),
    ("Mes a analizar", "Por defecto el Resumen muestra el mes actual. Para revisar otro mes, escribe cualquier fecha de ese mes en Configuración."),
]
r = 4
for a, b in rows:
    ws.cell(row=r, column=2, value=a)
    if b is None and a:
        ws.cell(row=r, column=2).font = font(bold=True, size=12, color=NAVY)
        ws.cell(row=r, column=2).border = Border(bottom=Side(style="medium", color=GOLD))
        ws.cell(row=r, column=3).border = Border(bottom=Side(style="medium", color=GOLD))
    else:
        ws.cell(row=r, column=2).font = font(bold=True)
        c = ws.cell(row=r, column=3, value=b); c.font = font()
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=2).alignment = Alignment(vertical="top")
    r += 1
ws["B16"].fill = fill(INPUT); ws["B16"].font = font(bold=True, color="0000FF")
ws["B17"].fill = fill(GRAY)

# ------------------------------------------------------------------ PRINCIPIOS
wp = wb.create_sheet("Principios")
title(wp, "PRINCIPIOS QUE GUÍAN ESTA PLANTILLA", "Inspirados en 'El hombre más rico de Babilonia' (G. S. Clason) y 'Padre Rico, Padre Pobre' (R. Kiyosaki)", "D")
widths(wp, {"A": 3, "B": 44, "C": 70, "D": 3})
header(wp, 4, 2, ["Principio", "Cómo lo aplica tu plantilla"])
principios = [
    ("Babilonia · 1ª cura: Empieza a engordar tu bolsa", "Págate primero: al menos 10% de cada ingreso va a Ahorro e Inversión ANTES de gastar."),
    ("Babilonia · 2ª cura: Controla tus gastos", "Presupuesto por categoría y por periodo (día/semana) con semáforo. Distingue necesidades de deseos."),
    ("Babilonia · 3ª cura: Haz que tu oro se multiplique", "El ahorro no se queda quieto: registra inversiones y el ingreso pasivo que generan."),
    ("Babilonia · 4ª cura: Protege tus tesoros de la pérdida", "Fondo de emergencia con meta en meses de gastos antes de inversiones de riesgo."),
    ("Babilonia · 5ª cura: Haz de tu vivienda una inversión rentable", "Categoría Vivienda con presupuesto; registra si tu vivienda te genera rentas."),
    ("Babilonia · 6ª cura: Asegura un ingreso para el futuro", "Índice de libertad financiera: ingreso pasivo ÷ gastos de vida."),
    ("Babilonia · 7ª cura: Aumenta tu habilidad para ganar", "Categoría 'Educación financiera' dentro de Ahorro e Inversión."),
    ("Babilonia · Las deudas: 20% para pagarlas", "Regla de Arkad adaptada: 10% ahorro, 20% deudas, 70% para vivir. El plan de deudas usa ese 20%."),
    ("Padre Rico · Los ricos compran activos", "Un activo pone dinero en tu bolsillo; un pasivo lo saca. Cada ingreso pasivo cuenta en tu tablero."),
    ("Padre Rico · Págate a ti primero", "La columna 'Real' de ahorro se compara contra la meta del mes."),
    ("Padre Rico · Aprende a leer tus números", "El Resumen es tu estado financiero personal: ingresos, gastos, activos y pasivos."),
    ("Padre Rico · Las deudas malas pagan los lujos", "Las tarjetas de crédito con intereses van primero: método avalancha (tasa más alta primero)."),
    ("Padre Rico · No trabajes por dinero, haz que el dinero trabaje para ti", "Meta final: que el ingreso pasivo cubra el 100% de tus gastos de vida."),
    ("Préstamos familiares sin intereses", "Se pagan en segundo plano, del saldo más pequeño al más grande (bola de nieve), para cumplir tu palabra sin descuidar las deudas caras."),
]
for i, (a, b) in enumerate(principios):
    rr = 5 + i
    wp.cell(row=rr, column=2, value=a).font = font(bold=True, color=NAVY)
    c = wp.cell(row=rr, column=3, value=b); c.font = font()
    for col in (2, 3):
        wp.cell(row=rr, column=col).alignment = Alignment(wrap_text=True, vertical="top")
        wp.cell(row=rr, column=col).border = BOX
    wp.row_dimensions[rr].height = 30

header(wp, 21, 2, ["Frase del día (se muestra en el Resumen)", "Fuente"])
frases = [
    ("Una parte de todo lo que ganas es tuya para conservarla.", "El hombre más rico de Babilonia"),
    ("Paga primero a ti mismo, después a los demás.", "Padre Rico, Padre Pobre"),
    ("Lo que llamamos 'gastos necesarios' siempre crece hasta igualar nuestros ingresos, a menos que lo impidamos.", "El hombre más rico de Babilonia"),
    ("Los ricos compran activos; los pobres y la clase media compran pasivos creyendo que son activos.", "Padre Rico, Padre Pobre"),
    ("Donde hay determinación, hay un camino.", "El hombre más rico de Babilonia"),
    ("No es cuánto dinero ganas, sino cuánto dinero conservas.", "Padre Rico, Padre Pobre"),
    ("Cada moneda ahorrada es un esclavo que trabajará para ti.", "El hombre más rico de Babilonia"),
    ("La educación financiera vale más que el dinero.", "Padre Rico, Padre Pobre"),
    ("Un hombre honrado paga sus deudas: destina una parte fija de tus ingresos a ellas.", "El hombre más rico de Babilonia"),
    ("Tu casa no es un activo si saca dinero de tu bolsillo.", "Padre Rico, Padre Pobre"),
    ("Protege tu tesoro: invierte solo donde el capital esté seguro.", "El hombre más rico de Babilonia"),
    ("El dinero va y viene; aprender cómo funciona es lo que te hace rico.", "Padre Rico, Padre Pobre"),
    ("La buena suerte llega a quien aprovecha las oportunidades.", "El hombre más rico de Babilonia"),
    ("Haz que el dinero trabaje para ti.", "Padre Rico, Padre Pobre"),
    ("Hoy es el mejor día para empezar a engordar tu bolsa.", "El hombre más rico de Babilonia"),
]
for i, (a, b) in enumerate(frases):
    rr = 22 + i
    c = wp.cell(row=rr, column=2, value=a); c.font = font(italic=True); c.alignment = Alignment(wrap_text=True)
    wp.cell(row=rr, column=3, value=b).font = font(color="7F7F7F")
    for col in (2, 3): wp.cell(row=rr, column=col).border = BOX
    wp.row_dimensions[rr].height = 28
wp["B38"] = "Frases parafraseadas de las obras citadas, con fines motivacionales."
wp["B38"].font = font(size=8, italic=True, color="7F7F7F")
name(wb, "Frases", "Principios!$B$22:$B$36")
name(wb, "Frases_Fuente", "Principios!$C$22:$C$36")

# ------------------------------------------------------------------ CONFIGURACION
wc = wb.create_sheet("Configuración")
title(wc, "CONFIGURACIÓN", "Llena solo las celdas amarillas", "I")
widths(wc, {"A": 38, "B": 18, "C": 16, "D": 3, "E": 22, "F": 22, "G": 24, "H": 16, "I": 12})
cfg = [
    (4, "Tu nombre", "Tu nombre", None, True),
    (5, "Frecuencia de control de gastos", "Semanal", None, True),
    (6, "Mes a analizar (opcional, vacío = mes actual)", None, DATEF, True),
    (7, "Mes en análisis (automático)", '=IF(B6="",DATE(YEAR(TODAY()),MONTH(TODAY()),1),DATE(YEAR(B6),MONTH(B6),1))', MONTHF, False),
    (8, "% Págate primero (ahorro e inversión)", 0.10, "0%", True),
    (9, "% Pago de deudas", 0.20, "0%", True),
    (10, "% Gastos de vida (automático)", "=1-B8-B9", "0%", False),
    (11, "Presupuesto fijo mensual para deudas (opcional)", None, MONEY, True),
    (12, "Meta del fondo de emergencia (meses de gastos)", 3, "0", True),
]
for rr, lab, val, fmt, inp in cfg:
    wc.cell(row=rr, column=1, value=lab).font = font(bold=True)
    wc.cell(row=rr, column=1).border = BOX
    c = wc.cell(row=rr, column=2, value=val)
    c.border = BOX
    c.font = font(color="0000FF") if inp else font()
    c.fill = fill(INPUT) if inp else fill(GRAY)
    if fmt: c.number_format = fmt
    c.alignment = Alignment(horizontal="center")
wc["C8"] = "Regla de Babilonia: 10%"; wc["C9"] = "Regla de Babilonia: 20%"; wc["C10"] = "Regla de Babilonia: 70%"
wc["C11"] = "Si lo dejas vacío se usa el % de deudas × ingreso del mes"
for a in ("C8", "C9", "C10", "C11"): wc[a].font = font(size=8, italic=True, color="7F7F7F")
wc["C10"].value = '=IF(B10<0,"⚠ Los % suman más de 100%","Regla de Babilonia: 70%")'
wc["B6"].comment = Comment("Escribe cualquier fecha del mes que quieras revisar (ej. 15/08/2026). Déjalo vacío para ver el mes actual.", "Plantilla")
wc["B11"].comment = Comment("Opcional. Si quieres destinar una cantidad fija a deudas en lugar del porcentaje, escríbela aquí.", "Plantilla")
name(wb, "Nombre", "'Configuración'!$B$4")
name(wb, "Frecuencia", "'Configuración'!$B$5")
name(wb, "Mes_Analisis", "'Configuración'!$B$7")
name(wb, "Pct_Ahorro", "'Configuración'!$B$8")
name(wb, "Pct_Deuda", "'Configuración'!$B$9")
name(wb, "Pct_Vida", "'Configuración'!$B$10")
name(wb, "Deuda_Fija", "'Configuración'!$B$11")
name(wb, "Meses_Emergencia", "'Configuración'!$B$12")

header(wc, 15, 1, ["Categoría de gasto", "Tipo", "Presupuesto mensual"])
cats = [("Vivienda / Renta", "Necesidad", 6500), ("Servicios (luz, agua, gas)", "Necesidad", 900),
        ("Súper / Despensa", "Necesidad", 3800), ("Transporte / Gasolina", "Necesidad", 1800),
        ("Teléfono / Internet", "Necesidad", 650), ("Salud / Farmacia", "Necesidad", 500),
        ("Educación", "Necesidad", 400), ("Mascotas", "Necesidad", 300),
        ("Comida fuera / Antojos", "Deseo", 1200), ("Entretenimiento", "Deseo", 600),
        ("Ropa / Calzado", "Deseo", 500), ("Suscripciones", "Deseo", 350),
        ("Regalos", "Deseo", 300), ("Otros", "Deseo", 400)]
for i in range(20):
    rr = 16 + i
    if i < len(cats):
        for j, v in enumerate(cats[i]): wc.cell(row=rr, column=1 + j, value=v)
style_range(wc, "A16:B35", inp=True)
style_range(wc, "C16:C35", MONEY, inp=True)
wc["A36"] = "Total presupuesto de categorías"; wc["A36"].font = font(bold=True)
wc["C36"] = "=SUM(C16:C35)"; wc["C36"].number_format = MONEY; wc["C36"].font = font(bold=True); wc["C36"].fill = fill(GRAY)
dv_list(wc, '"Necesidad,Deseo"', "B16:B35")
name(wb, "Lista_Categorias", "'Configuración'!$A$16:$A$35")
name(wb, "Cat_Tipo", "'Configuración'!$B$16:$B$35")
name(wb, "Cat_Presupuesto", "'Configuración'!$C$16:$C$35")

lists = {
    "E": ("Fuentes de ingreso", ["Sueldo", "Negocio", "Freelance / Extra", "Ingreso pasivo", "Otro"]),
    "F": ("Tipos de deuda", ["Tarjeta de crédito", "Préstamo personal", "Crédito auto", "Crédito hipotecario", "Préstamo familiar", "Préstamo amigo", "Otro"]),
    "G": ("Destinos de ahorro", ["Fondo de emergencia", "Inversión (bolsa, CETES, fondos)", "Negocio", "Bienes raíces", "Educación financiera", "Retiro / Afore", "Otro"]),
    "H": ("Métodos de pago", ["Efectivo", "Débito", "Transferencia", "Tarjeta de crédito"]),
    "I": ("Frecuencias", ["Diario", "Semanal", "Mensual"]),
}
for col, (hdr, vals) in lists.items():
    c = wc[f"{col}15"]; c.value = hdr; c.font = font(bold=True, color="FFFFFF"); c.fill = fill(GOLD); c.border = BOX
    c.alignment = Alignment(horizontal="center", wrap_text=True)
    for i in range(10):
        cc = wc[f"{col}{16+i}"]
        if i < len(vals): cc.value = vals[i]
        cc.border = BOX; cc.font = font(color="0000FF"); cc.fill = fill(INPUT)
wc["E27"] = "Puedes agregar opciones en las celdas amarillas vacías de cada lista."
wc["E27"].font = font(size=8, italic=True, color="7F7F7F")
name(wb, "Lista_Fuentes", "'Configuración'!$E$16:$E$25")
name(wb, "Lista_TiposDeuda", "'Configuración'!$F$16:$F$25")
name(wb, "Lista_Destinos", "'Configuración'!$G$16:$G$25")
name(wb, "Lista_Metodos", "'Configuración'!$H$16:$H$25")
dv_list(wc, "$I$16:$I$18", "B5", "Elige Diario, Semanal o Mensual")

# ------------------------------------------------------------------ INGRESOS
N_ING = 300
wi = wb.create_sheet("Ingresos")
title(wi, "INGRESOS", "Sueldo, ganancias del negocio e ingreso pasivo. Todo lo que entra a tu bolsa.", "F")
widths(wi, {"A": 13, "B": 22, "C": 36, "D": 16, "E": 14, "F": 18})
header(wi, 4, 1, ["Fecha", "Fuente", "Descripción", "Monto", "Mes", "¿Trabajo o activo?"])
ex_ing = [(date(2026, 9, 1), "Sueldo", "Quincena 1", 11500), (date(2026, 9, 10), "Ingreso pasivo", "Rendimiento CETES", 280),
          (date(2026, 9, 15), "Sueldo", "Quincena 2", 11500), (date(2026, 9, 20), "Negocio", "Venta de postres", 3400)]
e1, eN = 5, 4 + N_ING
for i in range(N_ING):
    rr = 5 + i
    if i < len(ex_ing):
        for j, v in enumerate(ex_ing[i]): wi.cell(row=rr, column=1 + j, value=v)
    wi.cell(row=rr, column=5, value=f'=IF(A{rr}="","",DATE(YEAR(A{rr}),MONTH(A{rr}),1))')
    wi.cell(row=rr, column=6, value=f'=IF(B{rr}="","",IF(B{rr}="Ingreso pasivo","Activo (pasivo)","Trabajo (activo)"))')
style_range(wi, f"A5:A{eN}", DATEF, inp=True)
style_range(wi, f"B5:C{eN}", inp=True)
style_range(wi, f"D5:D{eN}", MONEY, inp=True)
style_range(wi, f"E5:E{eN}", "mmm yyyy", calc=True, align="center")
style_range(wi, f"F5:F{eN}", calc=True)
dv_list(wi, "Lista_Fuentes", f"B5:B{eN}")
wi.freeze_panes = "A5"; wi.auto_filter.ref = f"A4:F{eN}"
ING = lambda c: f"Ingresos!${c}$5:${c}${eN}"

# ------------------------------------------------------------------ GASTOS
N_G = 1000
wg = wb.create_sheet("Gastos")
title(wg, "REGISTRO DE GASTOS", "Anota cada gasto diario o semanal. 'Controla tus gastos' — 2ª cura de Babilonia.", "H")
widths(wg, {"A": 13, "B": 26, "C": 34, "D": 15, "E": 18, "F": 13, "G": 10, "H": 13})
header(wg, 4, 1, ["Fecha", "Categoría", "Descripción", "Monto", "Método de pago", "Necesidad / Deseo", "Semana", "Mes"])
ex_g = [(date(2026, 9, 1), "Vivienda / Renta", "Renta septiembre", 6500, "Transferencia"),
        (date(2026, 9, 2), "Súper / Despensa", "Despensa semanal", 920, "Débito"),
        (date(2026, 9, 3), "Transporte / Gasolina", "Gasolina", 600, "Débito"),
        (date(2026, 9, 5), "Comida fuera / Antojos", "Café y pan", 95, "Efectivo"),
        (date(2026, 9, 6), "Servicios (luz, agua, gas)", "Recibo de luz", 480, "Transferencia"),
        (date(2026, 9, 8), "Teléfono / Internet", "Internet casa", 499, "Débito"),
        (date(2026, 9, 9), "Súper / Despensa", "Despensa semanal", 875, "Débito"),
        (date(2026, 9, 12), "Entretenimiento", "Cine", 260, "Tarjeta de crédito"),
        (date(2026, 9, 14), "Salud / Farmacia", "Medicamento", 185, "Efectivo"),
        (date(2026, 9, 16), "Súper / Despensa", "Despensa semanal", 910, "Débito"),
        (date(2026, 9, 18), "Suscripciones", "Streaming", 219, "Tarjeta de crédito"),
        (date(2026, 9, 20), "Comida fuera / Antojos", "Comida familiar", 540, "Efectivo"),
        (date(2026, 9, 21), "Transporte / Gasolina", "Gasolina", 600, "Débito"),
        (date(2026, 9, 23), "Súper / Despensa", "Despensa semanal", 890, "Débito")]
gN = 4 + N_G
for i in range(N_G):
    rr = 5 + i
    if i < len(ex_g):
        for j, v in enumerate(ex_g[i]): wg.cell(row=rr, column=1 + j, value=v)
    wg.cell(row=rr, column=6, value=f'=IF(B{rr}="","",IFERROR(INDEX(Cat_Tipo,MATCH(B{rr},Lista_Categorias,0)),"Deseo"))')
    wg.cell(row=rr, column=7, value=f'=IF(A{rr}="","",WEEKNUM(A{rr},2))')
    wg.cell(row=rr, column=8, value=f'=IF(A{rr}="","",DATE(YEAR(A{rr}),MONTH(A{rr}),1))')
style_range(wg, f"A5:A{gN}", DATEF, inp=True)
style_range(wg, f"B5:C{gN}", inp=True)
style_range(wg, f"D5:D{gN}", MONEY, inp=True)
style_range(wg, f"E5:E{gN}", inp=True)
style_range(wg, f"F5:G{gN}", calc=True, align="center")
style_range(wg, f"H5:H{gN}", "mmm yyyy", calc=True, align="center")
dv_list(wg, "Lista_Categorias", f"B5:B{gN}", "Elige una categoría de la lista (agrégala en Configuración).")
dv_list(wg, "Lista_Metodos", f"E5:E{gN}")
wg.conditional_formatting.add(f"E5:E{gN}", CellIsRule(operator="equal", formula=['"Tarjeta de crédito"'], font=Font(name=F, color="C00000", bold=True)))
wg["J4"] = "Nota"; wg["J4"].font = font(bold=True, color=NAVY)
wg["J5"] = "Los gastos con tarjeta de crédito aparecen en rojo: si no liquidas el total, se convierten en deuda con intereses."
wg["J5"].font = font(size=9, italic=True, color="7F7F7F"); wg["J5"].alignment = Alignment(wrap_text=True, vertical="top")
wg.merge_cells("J5:M8"); widths(wg, {"J": 14})
wg.freeze_panes = "A5"; wg.auto_filter.ref = f"A4:H{gN}"
GAS = lambda c: f"Gastos!${c}$5:${c}${gN}"

# ------------------------------------------------------------------ ABONOS
N_P = 300
wa = wb.create_sheet("Abonos a Deudas")
title(wa, "ABONOS A DEUDAS", "Registra cada pago. El saldo de la deuda se actualiza solo en la hoja Deudas.", "E")
widths(wa, {"A": 13, "B": 30, "C": 15, "D": 34, "E": 13})
header(wa, 4, 1, ["Fecha", "Deuda", "Monto", "Nota", "Mes"])
ex_p = [(date(2026, 9, 5), "TDC Tienda departamental", 3150, "Mínimo + extra (prioridad 1)"),
        (date(2026, 9, 5), "TDC Banco Azul", 1200, "Pago mínimo"),
        (date(2026, 9, 10), "Préstamo personal", 1500, "Mensualidad"),
        (date(2026, 9, 15), "Préstamo hermano", 500, "Abono acordado")]
pN = 4 + N_P
for i in range(N_P):
    rr = 5 + i
    if i < len(ex_p):
        for j, v in enumerate(ex_p[i]): wa.cell(row=rr, column=1 + j, value=v)
    wa.cell(row=rr, column=5, value=f'=IF(A{rr}="","",DATE(YEAR(A{rr}),MONTH(A{rr}),1))')
style_range(wa, f"A5:A{pN}", DATEF, inp=True)
style_range(wa, f"B5:B{pN}", inp=True)
style_range(wa, f"C5:C{pN}", MONEY, inp=True)
style_range(wa, f"D5:D{pN}", inp=True)
style_range(wa, f"E5:E{pN}", "mmm yyyy", calc=True, align="center")
wa.freeze_panes = "A5"; wa.auto_filter.ref = f"A4:E{pN}"
PAG = lambda c: f"'Abonos a Deudas'!${c}$5:${c}${pN}"

# ------------------------------------------------------------------ AHORRO E INVERSION
N_A = 300
wv = wb.create_sheet("Ahorro e Inversión")
title(wv, "AHORRO E INVERSIÓN · PÁGATE PRIMERO", "'Una parte de todo lo que ganas es tuya para conservarla.' Registra aquí cada peso que pones a trabajar.", "E")
widths(wv, {"A": 13, "B": 32, "C": 15, "D": 34, "E": 13})
header(wv, 4, 1, ["Fecha", "Destino", "Monto", "Nota", "Mes"])
ex_a = [(date(2026, 9, 1), "Fondo de emergencia", 1500, "10% de la quincena"),
        (date(2026, 9, 15), "Inversión (bolsa, CETES, fondos)", 1000, "CETES 28 días"),
        (date(2026, 9, 20), "Educación financiera", 250, "Libro / curso")]
vN = 4 + N_A
for i in range(N_A):
    rr = 5 + i
    if i < len(ex_a):
        for j, v in enumerate(ex_a[i]): wv.cell(row=rr, column=1 + j, value=v)
    wv.cell(row=rr, column=5, value=f'=IF(A{rr}="","",DATE(YEAR(A{rr}),MONTH(A{rr}),1))')
style_range(wv, f"A5:A{vN}", DATEF, inp=True)
style_range(wv, f"B5:B{vN}", inp=True)
style_range(wv, f"C5:C{vN}", MONEY, inp=True)
style_range(wv, f"D5:D{vN}", inp=True)
style_range(wv, f"E5:E{vN}", "mmm yyyy", calc=True, align="center")
dv_list(wv, "Lista_Destinos", f"B5:B{vN}")
wv.freeze_panes = "A5"; wv.auto_filter.ref = f"A4:E{vN}"
AHO = lambda c: f"'Ahorro e Inversión'!${c}$5:${c}${vN}"

# ------------------------------------------------------------------ DEUDAS
wd = wb.create_sheet("Deudas", 3)
title(wd, "PLAN PARA ELIMINAR DEUDAS", "Prioridad 1: deudas con intereses (tasa más alta primero). Prioridad 2: préstamos familiares sin intereses (saldo más pequeño primero).", "M")
widths(wd, {"A": 28, "B": 20, "C": 15, "D": 11, "E": 14, "F": 15, "G": 15, "H": 14, "I": 10, "J": 15, "K": 12, "L": 14, "M": 34, "N": 10, "O": 12, "P": 12})
d0, dN = 9, 23
blk = [(4, "Presupuesto del mes para deudas", '=IF(Deuda_Fija<>"",Deuda_Fija,Resumen!$B$6*Pct_Deuda)'),
       (5, "Suma de pagos mínimos", f"=SUM(O{d0}:O{dN})"),
       (6, "Dinero extra para la deuda prioritaria", "=MAX(0,C4-C5)")]
for rr, lab, f_ in blk:
    wd.merge_cells(f"A{rr}:B{rr}")
    wd.cell(row=rr, column=1, value=lab).font = font(bold=True)
    c = wd.cell(row=rr, column=3, value=f_); c.number_format = MONEY; c.fill = fill(GRAY); c.border = BOX; c.font = font(bold=True)
wd["E4"] = '=IF(C4<C5,"⚠ Tu presupuesto no alcanza para los pagos mínimos. Recorta tus gastos tipo Deseo o aumenta tus ingresos.",IF(C6>0,"✔ Aplica el extra de "&TEXT(C6,"$#,##0")&" a la deuda con prioridad 1: así ahorras más intereses.","✔ Cubres tus mínimos."))'
wd["E4"].font = font(bold=True, color="7F6000"); wd.merge_cells("E4:M5"); wd["E4"].alignment = Alignment(wrap_text=True, vertical="center")
header(wd, 8, 1, ["Deuda", "Tipo", "Saldo inicial", "Tasa anual", "Pago mínimo mensual", "Abonos realizados",
                  "Saldo actual", "Interés mensual estimado", "Prioridad", "Pago sugerido este mes",
                  "Meses para liquidar", "Fecha estimada libre", "Estrategia", "Puntaje", "Mínimo aplicable", "Remanente"])
ex_d = [("TDC Banco Azul", "Tarjeta de crédito", 18500, 0.62, 1200),
        ("TDC Tienda departamental", "Tarjeta de crédito", 6200, 0.85, 650),
        ("Préstamo personal", "Préstamo personal", 25000, 0.28, 1500),
        ("Préstamo mamá", "Préstamo familiar", 10000, 0, 0),
        ("Préstamo hermano", "Préstamo familiar", 4000, 0, 500)]
for i in range(dN - d0 + 1):
    rr = d0 + i
    if i < len(ex_d):
        for j, v in enumerate(ex_d[i]): wd.cell(row=rr, column=1 + j, value=v)
    wd[f"F{rr}"] = f'=IF(A{rr}="",0,SUMIF({PAG("B")},A{rr},{PAG("C")}))'
    wd[f"G{rr}"] = f'=IF(A{rr}="","",MAX(0,C{rr}-F{rr}))'
    wd[f"H{rr}"] = f'=IF(A{rr}="","",G{rr}*D{rr}/12)'
    # Puntaje: con interés = 1 + tasa (más alta primero); sin interés = 1/(1+saldo) (más pequeña primero)
    wd[f"N{rr}"] = f'=IF(A{rr}="","",IF(G{rr}<=0,-1,IF(D{rr}>0,1+D{rr},1/(1+G{rr}))))'
    wd[f"I{rr}"] = f'=IF(OR(A{rr}="",N(G{rr})<=0),"",COUNTIF($N${d0}:$N${dN},">"&N{rr})+COUNTIF($N${d0}:N{rr},N{rr}))'
    wd[f"O{rr}"] = f'=IF(OR(A{rr}="",N(G{rr})<=0),0,MIN(G{rr},E{rr}))'
    wd[f"P{rr}"] = f'=IF(OR(A{rr}="",N(G{rr})<=0),0,G{rr}-O{rr})'
    wd[f"J{rr}"] = f'=IF(I{rr}="",0,O{rr}+MIN(P{rr},MAX(0,$C$6-SUMIF($I${d0}:$I${dN},"<"&I{rr},$P${d0}:$P${dN}))))'
    wd[f"K{rr}"] = (f'=IF(A{rr}="","",IF(N(G{rr})<=0,0,IF(J{rr}<=0,"Después de las anteriores",IF(D{rr}=0,ROUNDUP(G{rr}/J{rr},0),'
                    f'IF(J{rr}<=H{rr},"Nunca: pago < interés",ROUNDUP(NPER(D{rr}/12,-J{rr},G{rr}),0))))))')
    wd[f"L{rr}"] = f'=IF(ISNUMBER(K{rr}),EDATE(Mes_Analisis,K{rr}),"")'
    wd[f"M{rr}"] = (f'=IF(A{rr}="","",IF(N(G{rr})<=0,"✅ ¡Liquidada! Celébralo",IF(I{rr}=1,"🎯 Ataca aquí: todo el dinero extra",'
                    f'IF(D{rr}>0,"Paga el mínimo: genera intereses","Sin intereses: en segundo plano"))))')
style_range(wd, f"A{d0}:B{dN}", inp=True)
style_range(wd, f"C{d0}:C{dN}", MONEY, inp=True)
style_range(wd, f"D{d0}:D{dN}", "0.0%", inp=True)
style_range(wd, f"E{d0}:E{dN}", MONEY, inp=True)
style_range(wd, f"F{d0}:H{dN}", MONEY, calc=True)
style_range(wd, f"I{d0}:I{dN}", "0", calc=True, align="center")
style_range(wd, f"J{d0}:J{dN}", MONEY, calc=True)
style_range(wd, f"K{d0}:K{dN}", "0", calc=True, align="center")
style_range(wd, f"L{d0}:L{dN}", "mmm yyyy", calc=True, align="center")
style_range(wd, f"M{d0}:M{dN}", calc=True)
style_range(wd, f"N{d0}:N{dN}", "0.0000", calc=True)
style_range(wd, f"O{d0}:P{dN}", MONEY, calc=True)
for col in "NOP": wd.column_dimensions[col].hidden = True
dv_list(wd, "Lista_TiposDeuda", f"B{d0}:B{dN}")
wd[f"D{d0}"].comment = Comment("Tasa ANUAL (CAT o tasa de interés) como porcentaje: 62% para una tarjeta típica. Préstamos familiares: 0%.", "Plantilla")
tR = dN + 1
wd[f"A{tR}"] = "TOTAL"; wd[f"A{tR}"].font = font(bold=True, color="FFFFFF")
for col in "CEFGHJ":
    wd[f"{col}{tR}"] = f"=SUM({col}{d0}:{col}{dN})"
    wd[f"{col}{tR}"].number_format = MONEY
for col in "ABCDEFGHIJKLM":
    c = wd[f"{col}{tR}"]; c.fill = fill(NAVY); c.border = BOX
    if col != "A": c.font = font(bold=True, color="FFFFFF")
wd[f"K{tR}"] = f'=IF(COUNT(L{d0}:L{dN})=0,"",MAX(L{d0}:L{dN}))'
wd[f"K{tR}"].number_format = "mmm yyyy"; wd.merge_cells(f"K{tR}:L{tR}")
wd[f"M{tR}"] = "← Libre de deudas (estimado)"
wd.conditional_formatting.add(f"A{d0}:M{dN}", FormulaRule(formula=[f"$I{d0}=1"], fill=fill("FCE4D6"), font=Font(name=F, bold=True)))
wd.conditional_formatting.add(f"A{d0}:M{dN}", FormulaRule(formula=[f'AND($A{d0}<>"",N($G{d0})<=0)'], fill=fill("E2EFDA")))
wd[f"A{tR+2}"] = "Cómo funciona"; wd[f"A{tR+2}"].font = font(bold=True, color=NAVY)
notas = ["• Cada mes cubres el pago mínimo de TODAS tus deudas para proteger tu historial.",
         "• Todo el dinero extra va a la deuda con prioridad 1. Al liquidarla, su pago pasa a la siguiente (efecto avalancha).",
         "• Las deudas con intereses van primero, de la tasa más alta a la más baja: cada peso ahí te ahorra más intereses.",
         "• Los préstamos familiares sin intereses van después, del saldo más pequeño al más grande, para liquidarlos uno a uno.",
         "• 'Meses para liquidar' y 'Fecha estimada libre' suponen que repites el pago sugerido de este mes; al acelerar con el efecto avalancha terminarás antes."]
for i, t in enumerate(notas):
    wd[f"A{tR+3+i}"] = t; wd[f"A{tR+3+i}"].font = font(size=9)
wd.freeze_panes = "B9"
name(wb, "Lista_Deudas", f"Deudas!$A${d0}:$A${dN}")
dv_list(wa, "Lista_Deudas", f"B5:B{pN}", "Elige una deuda registrada en la hoja Deudas.")
DEU = lambda c: f"Deudas!${c}${d0}:${c}${dN}"

# ------------------------------------------------------------------ RESUMEN
wr = wb.create_sheet("Resumen", 1)
title(wr, "RESUMEN DEL MES", None, "Q")
wr["A1"] = '="TABLERO DE LIBERTAD FINANCIERA · "&UPPER(TEXT(Mes_Analisis,"mmmm yyyy"))&IF(OR(Nombre="",Nombre="Tu nombre"),""," · "&Nombre)'
wr["A2"] = '="💬 “"&INDEX(Frases,MOD(DAY(TODAY())-1,ROWS(Frases))+1)&"” — "&INDEX(Frases_Fuente,MOD(DAY(TODAY())-1,ROWS(Frases))+1)'
wr["A2"].font = font(italic=True, color="7F6000"); wr.merge_cells("A2:Q2"); wr["A2"].fill = fill(CREAM)
widths(wr, {"A": 2, "B": 30, "C": 16, "D": 16, "E": 16, "F": 16, "G": 24, "H": 3, "I": 34, "J": 16})

M = "Mes_Analisis"
kpis = [("B", "Ingresos del mes", f"=SUMIFS({ING('D')},{ING('E')},{M})"),
        ("C", "Gastos de vida", f"=SUMIFS({GAS('D')},{GAS('H')},{M})"),
        ("D", "Abonos a deudas", f"=SUMIFS({PAG('C')},{PAG('E')},{M})"),
        ("E", "Ahorro e inversión", f"=SUMIFS({AHO('C')},{AHO('E')},{M})"),
        ("F", "Disponible", "=B6-C6-D6-E6")]
for col, lab, f_ in kpis:
    wr[f"{col}5"] = lab; wr[f"{col}5"].font = font(bold=True, color="FFFFFF", size=9)
    wr[f"{col}5"].fill = fill(GOLD if col == "F" else NAVY); wr[f"{col}5"].alignment = Alignment(horizontal="center", wrap_text=True)
    wr[f"{col}6"] = f_; wr[f"{col}6"].font = font(bold=True, size=14, color=NAVY)
    wr[f"{col}6"].number_format = MONEY; wr[f"{col}6"].alignment = Alignment(horizontal="center")
    wr[f"{col}6"].fill = fill(CREAM); wr[f"{col}6"].border = BOX
wr.row_dimensions[6].height = 28
wr["G5"] = "Mensaje para ti"; wr["G5"].font = font(bold=True, color=NAVY)
wr["G6"] = ('=IF(B6=0,"Registra tus ingresos del mes para empezar.",IF(F6<0,"⚠ Gastas más de lo que ganas. Controla tus gastos (2ª cura de Babilonia).",'
            'IF(E6<B6*Pct_Ahorro,"Aún no te pagas primero: aparta "&TEXT(B6*Pct_Ahorro-E6,"$#,##0")&" para tu bolsa.",'
            '"¡Excelente! Te pagaste primero. Pon a trabajar tu oro.")))')
wr["G6"].font = font(bold=True, color="C00000"); wr["G6"].alignment = Alignment(wrap_text=True, vertical="center")
wr.merge_cells("G6:J6"); wr.merge_cells("G5:J5")
wr.conditional_formatting.add("F6", CellIsRule(operator="lessThan", formula=["0"], fill=fill("F8CBAD")))

# Regla de Babilonia
wr["B8"] = "REGLA DE BABILONIA (10 / 20 / 70)"; wr["B8"].font = font(bold=True, size=12, color=NAVY)
header(wr, 9, 2, ["Concepto", "% del ingreso", "Presupuesto", "Real", "Diferencia", "Estado"])
bab = [(10, "Págate primero (ahorro e inversión)", "=Pct_Ahorro", "=E6", "min"),
       (11, "Pago de deudas", "=Pct_Deuda", "=D6", "min"),
       (12, "Gastos de vida", "=Pct_Vida", "=C6", "max")]
for rr, lab, pct, real, kind in bab:
    wr[f"B{rr}"] = lab; wr[f"C{rr}"] = pct
    wr[f"D{rr}"] = "=Deudas!$C$4" if rr == 11 else f"=$B$6*C{rr}"
    wr[f"E{rr}"] = real; wr[f"F{rr}"] = f"=D{rr}-E{rr}"
    if kind == "min":
        wr[f"G{rr}"] = f'=IF(D{rr}=0,"-",IF(E{rr}>=D{rr},"✅ Meta cumplida","⏳ Faltan "&TEXT(D{rr}-E{rr},"$#,##0")))'
    else:
        wr[f"G{rr}"] = f'=IF(D{rr}=0,"-",IF(E{rr}<=D{rr},"✅ Dentro del presupuesto","⚠ Excedido por "&TEXT(E{rr}-D{rr},"$#,##0")))'
wr["B13"] = "Total"; wr["C13"] = "=SUM(C10:C12)"; wr["D13"] = "=SUM(D10:D12)"; wr["E13"] = "=SUM(E10:E12)"; wr["F13"] = "=D13-E13"
style_range(wr, "B10:B13"); style_range(wr, "C10:C13", "0%", align="center")
style_range(wr, "D10:F13", MONEY); style_range(wr, "G10:G13")
for col in "BCDEFG": wr[f"{col}13"].font = font(bold=True); wr[f"{col}13"].fill = fill(CREAM)
wr.conditional_formatting.add("G10:G12", FormulaRule(formula=['LEFT(G10,1)="⚠"'], font=Font(name=F, color="C00000", bold=True)))
wr.conditional_formatting.add("G10:G12", FormulaRule(formula=['LEFT(G10,1)="✅"'], font=Font(name=F, color="00803C", bold=True)))

# Control por periodo
wr["B15"] = "CONTROL DE GASTO POR PERIODO"; wr["B15"].font = font(bold=True, size=12, color=NAVY)
per = [(16, "Frecuencia elegida", "=Frecuencia", None),
       (17, "Presupuesto de vida del mes", "=D12", MONEY),
       (18, "Periodos en el mes", f'=IF(Frecuencia="Diario",DAY(EOMONTH({M},0)),IF(Frecuencia="Semanal",DAY(EOMONTH({M},0))/7,1))', "0.0"),
       (19, "Presupuesto por periodo", "=IF(C18=0,0,C17/C18)", MONEY),
       (20, "Periodo actual (hoy)", '=IF(Frecuencia="Diario",TEXT(TODAY(),"dd/mm/yyyy"),IF(Frecuencia="Semanal",TEXT(TODAY()-WEEKDAY(TODAY(),2)+1,"dd/mm")&" al "&TEXT(TODAY()-WEEKDAY(TODAY(),2)+7,"dd/mm"),TEXT(TODAY(),"mmmm yyyy")))', None),
       (21, "Gastado en el periodo actual", (f'=IF(Frecuencia="Diario",SUMIFS({GAS("D")},{GAS("A")},TODAY()),'
                                              f'IF(Frecuencia="Semanal",SUMIFS({GAS("D")},{GAS("A")},">="&(TODAY()-WEEKDAY(TODAY(),2)+1),{GAS("A")},"<="&(TODAY()-WEEKDAY(TODAY(),2)+7)),'
                                              f'SUMIFS({GAS("D")},{GAS("H")},DATE(YEAR(TODAY()),MONTH(TODAY()),1))))'), MONEY),
       (22, "Disponible en el periodo", "=C19-C21", MONEY),
       (23, "Semáforo", '=IF(C19=0,"-",IF(C21<=C19*0.8,"🟢 Vas bien",IF(C21<=C19,"🟡 Cuidado, cerca del límite","🔴 Te pasaste: frena los deseos")))', None)]
for rr, lab, f_, fmt in per:
    wr[f"B{rr}"] = lab; wr[f"C{rr}"] = f_
    wr.merge_cells(f"C{rr}:D{rr}")
    style_range(wr, f"B{rr}:D{rr}", fmt, calc=True)
    wr[f"B{rr}"].font = font(bold=True); wr[f"B{rr}"].fill = PatternFill()
    wr[f"C{rr}"].alignment = Alignment(horizontal="center")
wr.conditional_formatting.add("C22", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=F, color="C00000", bold=True)))

# Indicadores Padre Rico
wr["F15"] = "INDICADORES PADRE RICO"; wr["F15"].font = font(bold=True, size=12, color=NAVY)
ind = [(16, "Tasa de ahorro del mes", "=IF(B6=0,0,E6/B6)", PCT),
       (17, "Ingreso pasivo del mes (activos)", f'=SUMIFS({ING("D")},{ING("E")},{M},{ING("B")},"Ingreso pasivo")', MONEY),
       (18, "Índice de libertad financiera", "=IF(C6=0,0,G17/C6)", PCT),
       (19, "Fondo de emergencia acumulado", f'=SUMIFS({AHO("C")},{AHO("B")},"Fondo de emergencia",{AHO("A")},"<="&EOMONTH({M},0))', MONEY),
       (20, "Meta del fondo de emergencia", "=Meses_Emergencia*MAX(C6,D12)", MONEY),
       (21, "Avance del fondo", "=IF(G20=0,0,MIN(1,G19/G20))", PCT),
       (22, "Total invertido en activos", f'=SUMIFS({AHO("C")},{AHO("A")},"<="&EOMONTH({M},0))-G19', MONEY),
       (23, "Deuda total actual", f"=Deudas!$G${tR}", MONEY),
       (24, "Deuda con intereses (prioridad)", f'=SUMIFS({DEU("G")},{DEU("D")},">0")', MONEY),
       (25, "Deuda sin intereses (familiar)", "=G23-G24", MONEY),
       (26, "Intereses que pagas al mes (estimado)", f"=Deudas!$H${tR}", MONEY),
       (27, "Libre de deudas (estimado)", f'=IF(G23=0,"¡Ya eres libre!",Deudas!$K${tR})', "mmmm yyyy")]
for rr, lab, f_, fmt in ind:
    wr[f"F{rr}"] = lab; wr[f"G{rr}"] = f_
    style_range(wr, f"F{rr}:G{rr}", fmt, calc=True)
    wr[f"F{rr}"].font = font(bold=True); wr[f"F{rr}"].fill = PatternFill()
    wr[f"G{rr}"].alignment = Alignment(horizontal="center")
widths(wr, {"F": 34})
wr.conditional_formatting.add("G21", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=GOLD))
wr.conditional_formatting.add("G18", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="70AD47"))
wr["F28"] = "Índice de libertad = ingreso pasivo ÷ gastos de vida. Al llegar a 100% tus activos pagan tu vida."
wr["F28"].font = font(size=8, italic=True, color="7F7F7F")

# Gastos por categoría
wr["B30"] = "GASTOS POR CATEGORÍA DEL MES"; wr["B30"].font = font(bold=True, size=12, color=NAVY)
header(wr, 31, 2, ["Categoría", "Tipo", "Presupuesto", "Real", "Disponible", "% usado"])
c0, cN = 32, 51
for i in range(20):
    rr = c0 + i; src = 16 + i
    wr[f"B{rr}"] = f"=IF('Configuración'!A{src}=\"\",\"\",'Configuración'!A{src})"
    wr[f"C{rr}"] = f"=IF(B{rr}=\"\",\"\",'Configuración'!B{src})"
    wr[f"D{rr}"] = f"=IF(B{rr}=\"\",\"\",N('Configuración'!C{src}))"
    wr[f"E{rr}"] = f'=IF(B{rr}="","",SUMIFS({GAS("D")},{GAS("B")},B{rr},{GAS("H")},{M}))'
    wr[f"F{rr}"] = f'=IF(B{rr}="","",D{rr}-E{rr})'
    wr[f"G{rr}"] = f'=IF(B{rr}="","",IF(D{rr}=0,IF(E{rr}>0,1,0),E{rr}/D{rr}))'
style_range(wr, f"B{c0}:C{cN}"); style_range(wr, f"D{c0}:F{cN}", MONEY); style_range(wr, f"G{c0}:G{cN}", "0%", align="center")
wr.conditional_formatting.add(f"G{c0}:G{cN}", CellIsRule(operator="greaterThan", formula=["1"], fill=fill("F8CBAD"), font=Font(name=F, color="9C0006", bold=True)))
wr.conditional_formatting.add(f"G{c0}:G{cN}", CellIsRule(operator="between", formula=["0.8", "1"], fill=fill("FFEB9C")))
wr.conditional_formatting.add(f"F{c0}:F{cN}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=F, color="C00000", bold=True)))
wr[f"B{cN+1}"] = "Total"; wr[f"D{cN+1}"] = f"=SUM(D{c0}:D{cN})"; wr[f"E{cN+1}"] = f"=SUM(E{c0}:E{cN})"; wr[f"F{cN+1}"] = f"=D{cN+1}-E{cN+1}"
wr[f"G{cN+1}"] = f"=IF(D{cN+1}=0,0,E{cN+1}/D{cN+1})"
style_range(wr, f"B{cN+1}:C{cN+1}"); style_range(wr, f"D{cN+1}:F{cN+1}", MONEY); style_range(wr, f"G{cN+1}:G{cN+1}", "0%", align="center")
for col in "BCDEFG": wr[f"{col}{cN+1}"].font = font(bold=True); wr[f"{col}{cN+1}"].fill = fill(CREAM)
wr[f"E{cN+2}"] = f'=IF(ROUND(E{cN+1},2)<>ROUND(C6,2),"⚠ Hay gastos con categorías que no están en Configuración","")'
wr[f"E{cN+2}"].font = font(size=8, italic=True, color="C00000")

nd = cN + 4
wr[f"B{nd}"] = "NECESIDADES VS DESEOS"; wr[f"B{nd}"].font = font(bold=True, size=12, color=NAVY)
header(wr, nd + 1, 2, ["Tipo", "Monto", "% del gasto"])
for i, t in enumerate(["Necesidad", "Deseo"]):
    rr = nd + 2 + i
    wr[f"B{rr}"] = t
    wr[f"C{rr}"] = f'=SUMIFS({GAS("D")},{GAS("F")},"{t}",{GAS("H")},{M})'
    wr[f"D{rr}"] = f"=IF($C$6=0,0,C{rr}/$C$6)"
style_range(wr, f"B{nd+2}:B{nd+3}"); style_range(wr, f"C{nd+2}:C{nd+3}", MONEY); style_range(wr, f"D{nd+2}:D{nd+3}", "0%", align="center")
wr[f"B{nd+4}"] = "Tip: recorta primero los 'Deseos' para acelerar el pago de tus tarjetas."
wr[f"B{nd+4}"].font = font(size=8, italic=True, color="7F7F7F")

ch = BarChart(); ch.type = "col"; ch.title = "Regla de Babilonia: presupuesto vs real"
ch.add_data(Reference(wr, min_col=4, max_col=5, min_row=9, max_row=12), titles_from_data=True)
ch.set_categories(Reference(wr, min_col=2, min_row=10, max_row=12))
ch.height, ch.width = 7.5, 15; ch.y_axis.numFmt = '"$"#,##0'; ch.y_axis.majorGridlines = None
ch.series[0].graphicalProperties.solidFill = NAVY; ch.series[1].graphicalProperties.solidFill = GOLD
wr.add_chart(ch, "I8")
pie = PieChart(); pie.title = "Gasto real por categoría"
pie.add_data(Reference(wr, min_col=5, min_row=31, max_row=cN), titles_from_data=True)
pie.set_categories(Reference(wr, min_col=2, min_row=c0, max_row=cN))
pie.height, pie.width = 9, 15
pie.dataLabels = DataLabelList(); pie.dataLabels.showPercent = True
wr.add_chart(pie, "I30")
wr.freeze_panes = "A4"

# ------------------------------------------------------------------ TENDENCIA
wt = wb.create_sheet("Tendencia", 2)
title(wt, "TENDENCIA ANUAL", None, "I")
wt["A1"] = '="TENDENCIA "&YEAR(Mes_Analisis)'
widths(wt, {"A": 14, "B": 15, "C": 15, "D": 15, "E": 15, "F": 15, "G": 15, "H": 13, "I": 13})
header(wt, 4, 1, ["Mes", "Ingresos", "Gastos de vida", "Abonos a deudas", "Ahorro e inversión", "Disponible", "Ingreso pasivo", "Tasa de ahorro", "Índice libertad"])
for m in range(12):
    rr = 5 + m
    wt[f"A{rr}"] = f"=DATE(YEAR(Mes_Analisis),{m+1},1)"
    wt[f"B{rr}"] = f"=SUMIFS({ING('D')},{ING('E')},A{rr})"
    wt[f"C{rr}"] = f"=SUMIFS({GAS('D')},{GAS('H')},A{rr})"
    wt[f"D{rr}"] = f"=SUMIFS({PAG('C')},{PAG('E')},A{rr})"
    wt[f"E{rr}"] = f"=SUMIFS({AHO('C')},{AHO('E')},A{rr})"
    wt[f"F{rr}"] = f"=B{rr}-C{rr}-D{rr}-E{rr}"
    wt[f"G{rr}"] = f'=SUMIFS({ING("D")},{ING("E")},A{rr},{ING("B")},"Ingreso pasivo")'
    wt[f"H{rr}"] = f"=IF(B{rr}=0,0,E{rr}/B{rr})"
    wt[f"I{rr}"] = f"=IF(C{rr}=0,0,G{rr}/C{rr})"
style_range(wt, "A5:A16", "mmm yyyy", calc=True, align="center")
style_range(wt, "B5:G16", MONEY, calc=True)
style_range(wt, "H5:I16", PCT, calc=True, align="center")
wt["A17"] = "Total"
for col in "BCDEFG":
    wt[f"{col}17"] = f"=SUM({col}5:{col}16)"
wt["H17"] = "=IF(B17=0,0,E17/B17)"; wt["I17"] = "=IF(C17=0,0,G17/C17)"
style_range(wt, "A17:A17"); style_range(wt, "B17:G17", MONEY); style_range(wt, "H17:I17", PCT, align="center")
for col in "ABCDEFGHI": wt[f"{col}17"].font = font(bold=True); wt[f"{col}17"].fill = fill(CREAM)
lc = LineChart(); lc.title = "Ingresos, gastos y ahorro por mes"
lc.add_data(Reference(wt, min_col=2, max_col=5, min_row=4, max_row=16), titles_from_data=True)
lc.set_categories(Reference(wt, min_col=1, min_row=5, max_row=16))
lc.height, lc.width = 9, 24; lc.y_axis.numFmt = '"$"#,##0'; lc.x_axis.number_format = "mmm"
for s, col in zip(lc.series, [NAVY, "C00000", "ED7D31", GOLD]):
    s.graphicalProperties.line.solidFill = col; s.graphicalProperties.line.width = 28000; s.smooth = False
wt.add_chart(lc, "A20")

# Order sheets
order = ["Instrucciones", "Resumen", "Configuración", "Ingresos", "Gastos", "Deudas", "Abonos a Deudas", "Ahorro e Inversión", "Tendencia", "Principios"]
wb._sheets = [wb[n] for n in order]
tabs = {"Instrucciones": GOLD, "Resumen": NAVY, "Configuración": "7F7F7F", "Ingresos": "70AD47", "Gastos": "C00000",
        "Deudas": "ED7D31", "Abonos a Deudas": "ED7D31", "Ahorro e Inversión": GOLD, "Tendencia": NAVY, "Principios": GOLD}
for n, c in tabs.items(): wb[n].sheet_properties.tabColor = c
for w in wb.worksheets:
    w.page_setup.orientation = "landscape"; w.page_setup.fitToWidth = 1; w.page_setup.fitToHeight = 0
    w.sheet_properties.pageSetUpPr.fitToPage = True
wb.active = 1
wb.save(OUT)
print("ok")
